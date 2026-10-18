"""Egg drop detection, light advice, grown chicks, care calendar and the Excel export."""

from __future__ import annotations

import asyncio
from datetime import date, datetime, timedelta, timezone
import importlib.util
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from app.services.care import CareService, format_care_reminder
from app.services.eggs import EggService
from app.services.export import build_farm_workbook
from app.services.feeds import FeedService
from app.services.incubation import IncubationService
from app.services.insights import daylight, detect_egg_drop, grown_chick_groups, light_advice_lines
from app.services.poultry_advisor import PoultryAdvisorService
from app.services.stock import StockService
from app.storage.database import Database
from app.storage.repositories.batches import BatchRepository
from app.storage.repositories.eggs import EggRepository
from app.storage.repositories.feeds import FeedRepository
from app.storage.repositories.stock import StockRepository


TODAY = date(2026, 10, 17)


class FarmInsightsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database = Database(Path(self.temp_dir.name) / "t.db")
        self.database.initialize()
        self.feeds = FeedRepository(self.database)
        self.feed_service = FeedService(self.feeds)
        self.hens = self.feed_service.create_bird_group(user_id=1, name="Несушки", bird_count=12, role="hens")
        self.eggs = EggRepository(self.database)
        self.egg_service = EggService(self.eggs, self.feeds, timezone_name="Europe/Moscow")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _egg(self, day: date, count: int) -> None:
        self.eggs.create_entry(
            user_id=1, entry_date=day, eggs_count=count,
            active_hens_count=12, total_hens_count=12, excluded_hens_count=0,
        )

    def _history(self, usual: int, recent: int, *, skip_recent: int | None = None) -> None:
        for offset in range(4, 18):
            self._egg(TODAY - timedelta(days=offset), usual)
        for offset in (1, 2, 3):
            if offset != skip_recent:
                self._egg(TODAY - timedelta(days=offset), recent)

    def _advisor(self) -> PoultryAdvisorService:
        return PoultryAdvisorService(
            incubation_service=IncubationService(BatchRepository(self.database)),
            feed_service=self.feed_service,
            egg_service=self.egg_service,
            stock_service=StockService(StockRepository(self.database), self.feeds),
            timezone_name="Europe/Moscow",
        )

    # --- egg drop

    def test_drop_detected(self) -> None:
        self._history(10, 5)
        drop = detect_egg_drop(self.eggs, 1, today=TODAY)
        self.assertIsNotNone(drop)
        self.assertAlmostEqual(drop.drop_percent, 50)

    def test_no_drop_for_small_dip_or_missing_day(self) -> None:
        self._history(10, 8)
        self.assertIsNone(detect_egg_drop(self.eggs, 1, today=TODAY))

    def test_forgotten_entry_is_not_a_drop(self) -> None:
        self._history(10, 3, skip_recent=2)
        self.assertIsNone(detect_egg_drop(self.eggs, 1, today=TODAY))

    def test_alert_text_mentions_short_day(self) -> None:
        self._history(10, 4)
        advisor = self._advisor()
        drop = advisor.detect_egg_drop(1, today=TODAY)

        text = advisor.build_egg_drop_alert(1, drop, today=TODAY)

        self.assertIn("Яиц стало заметно меньше", text)
        self.assertIn("−60%", text)
        self.assertIn("Короткий день", text)

    # --- light

    def test_daylight_moscow(self) -> None:
        december = daylight(date(2026, 12, 22), latitude=55.75, longitude=37.62, timezone_name="Europe/Moscow")
        june = daylight(date(2026, 6, 21), latitude=55.75, longitude=37.62, timezone_name="Europe/Moscow")
        self.assertAlmostEqual(december.hours, 7.0, delta=0.2)
        self.assertAlmostEqual(june.hours, 17.5, delta=0.3)
        self.assertIn("не хватает", "\n".join(light_advice_lines(december)))
        self.assertIn("хватает", "\n".join(light_advice_lines(june)))

    def test_light_plan_never_starts_before_4_30(self) -> None:
        text = "\n".join(
            light_advice_lines(daylight(date(2026, 12, 22), latitude=55.75, longitude=37.62, timezone_name="Europe/Moscow"))
        )
        self.assertIn("свет с 04:30", text)
        self.assertIn("после заката", text)

    # --- chicks

    def test_grown_chicks_split_into_hens_and_roosters_in_same_flock(self) -> None:
        chicks = self.feed_service.create_bird_group(
            user_id=1, name="Молодняк", bird_count=10, group_kind="chicks", hatched_at=TODAY - timedelta(weeks=17)
        )
        young = self.feed_service.create_bird_group(
            user_id=1, name="Малыши", bird_count=5, group_kind="chicks", hatched_at=TODAY - timedelta(weeks=3)
        )
        flock = self.feed_service.create_flock(user_id=1, name="Стадо", member_group_ids=[self.hens.id, chicks.id])

        self.assertEqual(
            [g.id for g in grown_chick_groups(self.feed_service.list_bird_groups(1), today=TODAY)], [chicks.id]
        )
        created = self.feed_service.grow_chicks(user_id=1, group_id=chicks.id, hens=7)

        self.assertEqual([(g.role, g.bird_count) for g in created], [("hens", 7), ("roosters", 3)])
        self.assertFalse(self.feed_service.get_bird_group(chicks.id, 1).is_active)
        members = {m.bird_group_id for m in self.feed_service.list_flock_members(flock.id, 1)}
        self.assertTrue({g.id for g in created} <= members)
        self.assertEqual(self.egg_service.total_laying_hens(1), 19)
        with self.assertRaises(ValueError):
            self.feed_service.grow_chicks(user_id=1, group_id=young.id, hens=9)

    def test_unknown_sex_makes_one_mixed_group(self) -> None:
        chicks = self.feed_service.create_bird_group(
            user_id=1, name="Молодняк", bird_count=6, group_kind="chicks", hatched_at=TODAY - timedelta(weeks=16)
        )
        created = self.feed_service.grow_chicks(user_id=1, group_id=chicks.id, hens=None)
        self.assertEqual([(g.role, g.bird_count, g.group_kind) for g in created], [("mixed", 6, "adult")])

    # --- care

    def test_care_calendar(self) -> None:
        care = CareService(self.database)
        tasks = care.list(1, today=TODAY)
        self.assertEqual({t.code for t in tasks}, {"parasites", "deworming", "litter", "disinfection"})
        self.assertEqual(len(care.due_for_reminder(1, today=TODAY)), 4)
        self.assertIn("первое напоминание", format_care_reminder(tasks, TODAY))

        litter = next(t for t in tasks if t.code == "litter")
        care.mark_done(1, litter.id, today=TODAY)
        care.mark_reminded(1, [t.id for t in tasks if t.id != litter.id], today=TODAY)

        self.assertEqual(care.due_for_reminder(1, today=TODAY), [])
        self.assertEqual(len(care.due_for_reminder(1, today=TODAY + timedelta(days=3))), 3)  # ignored -> again
        self.assertEqual(care.get(1, litter.id).next_due, TODAY + timedelta(days=14))

        care.set_interval(1, litter.id, 10)
        self.assertEqual(care.get(1, litter.id).next_due, TODAY + timedelta(days=10))
        care.toggle(1, litter.id, today=TODAY)
        self.assertFalse(care.get(1, litter.id).is_active)
        with self.assertRaises(ValueError):
            care.set_interval(1, litter.id, 0)
        self.assertIsNone(care.mark_done(2, litter.id, today=TODAY))  # other user's task

    # --- export

    def test_excel_export_has_all_sheets(self) -> None:
        from openpyxl import load_workbook

        self._egg(TODAY, 9)
        StockService(StockRepository(self.database), self.feeds).add_purchase(
            user_id=1, name="Кукуруза", kind="ingredient", amount_kg=25
        )
        CareService(self.database).list(1, today=TODAY)

        book = load_workbook(BytesIO(build_farm_workbook(self.database, 1)))

        self.assertEqual(
            book.sheetnames,
            ["Яйца", "Кормления", "Кормления по ингредиентам", "Склад", "Продажа и расход яиц", "Потери",
             "Поголовье", "Инкубация", "Уход", "Хранение яиц"],
        )
        self.assertEqual(book["Яйца"]["B2"].value, 9)
        self.assertEqual(book["Склад"]["B2"].value, "Кукуруза")
        self.assertEqual(book["Поголовье"]["B2"].value, "Несушки")
        self.assertEqual(book["Потери"]["A2"].value, "Записей пока нет")
        # Another user's data never leaks into the file.
        other = load_workbook(BytesIO(build_farm_workbook(self.database, 2)))
        self.assertEqual(other["Яйца"]["A2"].value, "Записей пока нет")


@unittest.skipUnless(importlib.util.find_spec("aiogram"), "aiogram is not installed")
class InsightRunnerTest(FarmInsightsTest):
    def _runner(self, bot):
        from app.services.reminders import ReminderRunner
        from app.storage.repositories.notifications import NotificationRepository
        from app.storage.repositories.users import UserRepository

        users = UserRepository(self.database)
        users.upsert(user_id=1, first_name="Никита")
        return ReminderRunner(
            bot=bot,
            incubation_service=IncubationService(BatchRepository(self.database)),
            feed_service=self.feed_service,
            egg_service=self.egg_service,
            poultry_advisor_service=self._advisor(),
            care_service=CareService(self.database),
            users=users,
            notifications=NotificationRepository(self.database),
            timezone="Europe/Moscow",
        )

    class Bot:
        def __init__(self) -> None:
            self.messages: list[str] = []

        async def send_message(self, user_id, text, reply_markup=None) -> None:
            self.messages.append(text)

    def test_drop_alert_once_a_week(self) -> None:
        self._history(10, 4)
        bot = self.Bot()
        runner = self._runner(bot)
        moment = datetime(2026, 10, 17, 7, 0, tzinfo=timezone.utc)  # 10:00 Moscow

        asyncio.run(runner._send_egg_drop_alerts(moment))
        asyncio.run(runner._send_egg_drop_alerts(moment + timedelta(hours=5)))

        self.assertEqual(len(bot.messages), 1)
        self.assertIn("Яиц стало заметно меньше", bot.messages[0])

    def test_care_and_chick_alerts(self) -> None:
        self.feed_service.create_bird_group(
            user_id=1, name="Молодняк", bird_count=10, group_kind="chicks", hatched_at=TODAY - timedelta(weeks=17)
        )
        bot = self.Bot()
        runner = self._runner(bot)
        moment = datetime(2026, 10, 17, 7, 0, tzinfo=timezone.utc)

        asyncio.run(runner._send_care_reminders(moment))
        asyncio.run(runner._send_care_reminders(moment + timedelta(hours=1)))
        asyncio.run(runner._send_grown_chick_alerts(moment))
        asyncio.run(runner._send_grown_chick_alerts(moment + timedelta(days=1)))

        self.assertEqual(sum("Пора по уходу" in m for m in bot.messages), 1)
        self.assertEqual(sum("подросли" in m for m in bot.messages), 1)


if __name__ == "__main__":
    unittest.main()
