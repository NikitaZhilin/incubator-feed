"""All farm data in one Excel workbook: eggs, feeding, warehouse, money, losses, flock, incubation, care."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from app.services.ledger import LOSS_REASONS, OUTFLOW_KINDS
from app.storage.database import Database


STOCK_TYPES = {
    "purchase": "Покупка",
    "mix_input": "В замес",
    "mix_output": "Замес готов",
    "manual_adjustment": "Корректировка",
    "write_off": "Списание",
}
STOCK_KINDS = {
    "ingredient": "Ингредиент",
    "finished_mix": "Готовая смесь",
    "commercial_feed": "Готовый корм",
    "other": "Другое",
}
ROLES = {"hens": "Несушки", "roosters": "Петухи", "chicks": "Цыплята", "mixed": "Смешанная"}
def _species() -> dict[str, str]:
    from app.domain import PROFILES

    return {code: profile.title for code, profile in PROFILES.items()}


SPECIES = _species()


def _zone(name: str) -> ZoneInfo:
    try:
        return ZoneInfo(name)
    except (ZoneInfoNotFoundError, ValueError):
        return ZoneInfo("Europe/Moscow")


def _local(value, zone: ZoneInfo):
    """SQLite timestamp text -> local naive datetime for Excel."""
    if value in (None, ""):
        return None
    text = str(value).replace("Z", "+00:00")
    try:
        moment = datetime.fromisoformat(text)
    except ValueError:
        try:
            moment = datetime.strptime(text[:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return text
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(zone).replace(tzinfo=None, microsecond=0)


def _day(value):
    if value in (None, ""):
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return value


def _sheets(database: Database, user_id: int, zone: ZoneInfo) -> list[tuple[str, list[str], list[list]]]:
    with database.connect() as c:
        def rows(sql: str):
            return c.execute(sql, (user_id,)).fetchall()

        eggs = [
            [_day(r["entry_date"]), r["eggs_count"], r["active_hens_count"], r["total_hens_count"], r["note"] or ""]
            for r in rows("SELECT * FROM egg_entries WHERE user_id = ? ORDER BY entry_date, id")
        ]
        feeding = [
            [_local(r["fed_at"], zone), r["cups"], r["mix_kg"], r["daily_norm_kg"], r["bird_count"], r["note"] or ""]
            for r in rows("SELECT * FROM feeding_events WHERE user_id = ? ORDER BY fed_at, id")
        ]
        feeding_items = [
            [_day(r["fed_date"]), r["name"], r["kg"], "да" if r["stock_item_id"] else "нет"]
            for r in rows(
                """
                SELECT e.fed_date, i.name, i.kg, i.stock_item_id
                FROM feeding_event_items AS i JOIN feeding_events AS e ON e.id = i.event_id
                WHERE e.user_id = ? ORDER BY e.fed_at, i.id
                """
            )
        ]
        stock = [
            [
                _local(r["occurred_at"] or r["created_at"], zone),
                r["name"],
                STOCK_KINDS.get(r["kind"], r["kind"]),
                STOCK_TYPES.get(r["type"], r["type"]),
                r["amount_kg"],
                r["balance_after_kg"],
                r["price_rub"],
                (r["price_rub"] / r["amount_kg"]) if r["price_rub"] and r["amount_kg"] else None,
                r["note"] or "",
            ]
            for r in rows(
                """
                SELECT t.*, i.name, i.kind FROM stock_transactions AS t
                JOIN stock_items AS i ON i.id = t.stock_item_id
                WHERE t.user_id = ? ORDER BY t.created_at, t.id
                """
            )
        ]
        outflows = [
            [_day(r["occurred_on"]), OUTFLOW_KINDS.get(r["kind"], r["kind"]), r["eggs_count"], r["amount_rub"], r["note"] or ""]
            for r in rows("SELECT * FROM egg_outflows WHERE user_id = ? ORDER BY occurred_on, id")
        ]
        losses = [
            [_day(r["occurred_on"]), r["group_name"], r["count"], LOSS_REASONS.get(r["reason"], r["reason"]), r["note"] or ""]
            for r in rows("SELECT * FROM bird_losses WHERE user_id = ? ORDER BY occurred_on, id")
        ]
        groups = [
            [
                r["name"],
                ROLES.get(r["role"], r["role"]),
                r["bird_count"],
                SPECIES.get(r["species"] or "", r["species"] or ""),
                _day(r["hatched_at"]),
                _day(r["joined_at"]),
                "да" if r["is_active"] else "архив",
            ]
            for r in rows("SELECT * FROM bird_groups WHERE user_id = ? ORDER BY is_active DESC, id")
        ]
        batches = [
            [
                r["title"] or "",
                SPECIES.get(r["species"], r["species"]),
                r["eggs_count"],
                _day(r["start_date"]),
                r["hatched_count"],
                (r["hatched_count"] / r["eggs_count"]) if r["hatched_count"] is not None and r["eggs_count"] else None,
                "идёт" if r["is_active"] else "завершена",
                r["infertile"],
                r["dead_embryo"],
                r["not_hatched"],
                r["note"] or "",
            ]
            for r in rows(
                """
                SELECT b.*, l.infertile, l.dead_embryo, l.not_hatched
                FROM incubation_batches AS b LEFT JOIN batch_losses AS l ON l.batch_id = b.id
                WHERE b.user_id = ? ORDER BY b.start_date, b.id
                """
            )
        ]
        care = [
            [r["title"], r["interval_days"], _day(r["last_done"]), _day(r["next_due"]), "да" if r["is_active"] else "нет"]
            for r in rows("SELECT * FROM care_tasks WHERE user_id = ? ORDER BY id")
        ]
    return [
        ("Яйца", ["Дата", "Яиц", "Несушек в расчёте", "Несушек всего", "Комментарий"], eggs),
        ("Кормления", ["Когда", "Кружек", "Смеси, кг", "Норма стада, кг/день", "Птиц", "Комментарий"], feeding),
        ("Кормления по ингредиентам", ["Дата", "Ингредиент", "Кг", "Списано со склада"], feeding_items),
        (
            "Склад",
            ["Когда", "Позиция", "Тип позиции", "Операция", "Кг", "Остаток после, кг", "Цена, ₽", "₽ за кг", "Комментарий"],
            stock,
        ),
        ("Продажа и расход яиц", ["Дата", "Куда", "Яиц", "Сумма, ₽", "Комментарий"], outflows),
        ("Потери", ["Дата", "Группа", "Птиц", "Причина", "Комментарий"], losses),
        ("Поголовье", ["Группа", "Роль", "Птиц", "Вид", "Дата вывода", "Дата подсадки", "Активна"], groups),
        (
            "Инкубация",
            ["Партия", "Вид", "Яиц заложено", "Дата закладки", "Вывелось", "Выводимость", "Статус",
             "Неоплод", "Замершие", "Не вылупились", "Комментарий"],
            batches,
        ),
        ("Уход", ["Задача", "Каждые, дн.", "Последний раз", "Следующий раз", "Включено"], care),
    ]


def build_farm_workbook(database: Database, user_id: int, *, timezone_name: str = "Europe/Moscow") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    zone = _zone(timezone_name)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F6B4F")
    for title, headers, data in _sheets(database, user_id, zone):
        sheet = workbook.create_sheet(title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        if data:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"
        for index, header in enumerate(headers, start=1):
            letter = get_column_letter(index)
            values = [header, *(row[index - 1] for row in data[:500])]
            width = max(len(str(value)) if value is not None else 0 for value in values)
            sheet.column_dimensions[letter].width = min(max(width + 2, 10), 45)
            for cell in sheet[letter][1:]:
                if isinstance(cell.value, datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"
                elif isinstance(cell.value, float):
                    cell.number_format = "0%" if header == "Выводимость" else "0.00"
        if not data:
            sheet.append(["Записей пока нет"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(today: date) -> str:
    return f"hozyaystvo_{today.isoformat()}.xlsx"
